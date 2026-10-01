"""Files end to end: the assistant generates them in a turn, the customer downloads them."""

from __future__ import annotations

import io
import json

from openpyxl import load_workbook

from tests.conftest import CUSTOMER_B
from tests.integration.test_agent import Harness
from tests.scripted_llm import ScriptedChatModel, answer, tool_call

STATEMENT = {
    "filename": "movimientos_junio",
    "format": "xlsx",
    "sheets": [
        {
            "name": "Movimientos",
            "columns": [
                {"header": "Fecha", "type": "date"},
                {"header": "Comercio", "type": "text"},
                {"header": "Monto", "type": "money"},
            ],
            "rows": [["2026-06-16", "Uber", 120]],
        }
    ],
}
SUMMARY = {
    "filename": "resumen",
    "format": "csv",
    "sheets": [
        {"name": "Resumen", "columns": [{"header": "Total", "type": "money"}], "rows": [[120]]}
    ],
}


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def test_assistant_generates_several_files_and_the_customer_downloads_them():
    def final(messages):
        payload = json.loads(messages[-1].content)
        assert payload["ok"] and len(payload["data"]["files"]) == 2
        return answer("Listo: te preparé el Excel con tus movimientos y un CSV con el total.")

    llm = ScriptedChatModel(
        script=[
            tool_call("search_transactions", {"merchant": "Uber"}),
            tool_call("generate_files", {"files": [STATEMENT, SUMMARY]}, call_id="call_2"),
            final,
        ]
    )
    with Harness(llm) as h:
        token = h.login()
        r = h.chat(token, "Quiero un Excel con mis pagos de Uber y un CSV con el total")
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["status"] == "answered"
        xlsx, csv = body["files"]
        assert xlsx["filename"] == "movimientos_junio.xlsx" and xlsx["format"] == "xlsx"
        assert csv["filename"] == "resumen.csv" and csv["rows"] == 1

        download = h.client.get(xlsx["download_url"], headers=_auth(token))
        assert download.status_code == 200
        assert download.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert 'filename="movimientos_junio.xlsx"' in download.headers["content-disposition"]
        assert download.headers["cache-control"] == "private, no-store"
        ws = load_workbook(io.BytesIO(download.content))["Movimientos"]
        assert ws["B2"].value == "Uber" and ws["C2"].value == 120

        text = h.client.get(csv["download_url"], headers=_auth(token))
        assert text.content.decode("utf-8-sig").splitlines() == ["Total", "120.00"]

        # The next turn starts without the previous turn's files.
        llm.script.append(answer("¿Algo más?"))
        again = h.chat(token, "gracias", body["conversation_id"]).json()
        assert again["files"] == []

        # The trace keeps the call, not the whole table.
        events = h.trace(token, body["conversation_id"]).json()["events"]
        [gen] = [e for e in events if e.get("tool") == "generate_files"]
        assert gen["outcome"] == "ok"


def test_files_are_only_for_their_customer_and_need_a_session():
    llm = ScriptedChatModel()
    with Harness(llm) as h:
        token = h.login()
        r = h.client.post("/v1/files", json={"files": [STATEMENT]}, headers=_auth(token))
        assert r.status_code == 201, r.text
        [ref] = r.json()["files"]

        other = h.login(CUSTOMER_B)
        assert h.client.get(ref["download_url"], headers=_auth(other)).status_code == 404
        assert h.client.get(ref["download_url"]).status_code == 401
        assert h.client.get("/v1/files/fil_unknown", headers=_auth(token)).status_code == 404
        assert h.client.get(ref["download_url"], headers=_auth(token)).status_code == 200


def test_posted_payload_is_validated_and_scoped_to_owned_conversations():
    with Harness(ScriptedChatModel()) as h:
        token = h.login()
        bad = {"files": [{**SUMMARY, "sheets": SUMMARY["sheets"] * 2}]}
        assert h.client.post("/v1/files", json=bad, headers=_auth(token)).status_code == 422
        foreign = {"conversation_id": "conv_" + "0" * 32, "files": [SUMMARY]}
        r = h.client.post("/v1/files", json=foreign, headers=_auth(token))
        assert r.status_code == 404 and r.json()["error"] == "conversation_not_found"
