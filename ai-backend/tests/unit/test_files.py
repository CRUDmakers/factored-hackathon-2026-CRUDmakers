"""Excel/CSV generation from the assistant's JSON: typing, safety, limits, the tool."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from ai_backend.files import generator, service
from ai_backend.files.models import ColumnSpec, FileSpec, GenerateFilesArgs
from ai_backend.files.store import MemoryFileStore, SqliteFileStore
from ai_backend.tools import definitions as d
from ai_backend.tools.registry import get_tool, parameters_schema
from tests.conftest import CUSTOMER_A, NOW

COLUMNS = [
    {"header": "Fecha", "type": "date"},
    {"header": "Comercio", "type": "text"},
    {"header": "Monto", "type": "money"},
    {"header": "Moneda"},
]
ROWS = [
    ["2026-06-15", "Uber", 12.5, "USD"],
    ["2026-06-16T10:30:00Z", '=HYPERLINK("http://x")', "1500", "MXN"],
    ["no es fecha", None, "abc", "COP"],
]


def _spec(fmt: str = "xlsx", **overrides) -> FileSpec:
    body = {
        "filename": "movimientos junio",
        "format": fmt,
        "sheets": [{"name": "Movimientos", "columns": COLUMNS, "rows": ROWS}],
    }
    return FileSpec.model_validate({**body, **overrides})


def test_xlsx_has_typed_cells_and_formulas_stay_text():
    content, rows = generator.build(_spec())
    assert rows == 3
    ws = load_workbook(io.BytesIO(content))["Movimientos"]
    assert [c.value for c in ws[1]] == ["Fecha", "Comercio", "Monto", "Moneda"]
    assert ws["A2"].value == datetime(2026, 6, 15) and ws["A2"].number_format == "yyyy-mm-dd"
    assert ws["A3"].value == datetime(2026, 6, 16, 10, 30)
    assert ws["C2"].value == 12.5 and ws["C2"].number_format == "#,##0.00"
    assert ws["C3"].value == 1500
    # Injected formula: written as text, never as a formula.
    assert ws["B3"].data_type == "s" and ws["B3"].value.startswith("=HYPERLINK")
    # What doesn't parse stays as the model wrote it.
    assert ws["A4"].value == "no es fecha" and ws["C4"].value == "abc" and ws["B4"].value is None
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D4"


def test_csv_is_utf8_with_bom_and_escapes_formulas():
    content, _ = generator.build(_spec("csv"))
    assert content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows[0] == ["Fecha", "Comercio", "Monto", "Moneda"]
    assert rows[1] == ["2026-06-15", "Uber", "12.50", "USD"]
    assert rows[2][0] == "2026-06-16 10:30" and rows[2][1].startswith("'=HYPERLINK")
    assert rows[3] == ["no es fecha", "", "abc", "COP"]


def test_typed_values():
    money = ColumnSpec(header="x", type="money")
    assert generator.typed("10", money) == Decimal("10.00")
    assert generator.typed("NaN", money) == "NaN"
    assert generator.typed("2026-01-02", ColumnSpec(header="d", type="date")) == date(2026, 1, 2)
    assert generator.typed(True, ColumnSpec(header="n", type="number")) is True


def test_download_name_is_safe_and_has_the_extension():
    assert _spec(filename='../../etc/pa"ss').download_name() == "etc_pa_ss.xlsx"
    assert _spec(filename="extracto.xlsx").download_name() == "extracto.xlsx"
    assert _spec("csv", filename="año 2026").download_name() == "año 2026.csv"


@pytest.mark.parametrize(
    "change, problem",
    [
        ({"format": "csv", "sheets": [{"name": "A", "columns": COLUMNS, "rows": []}] * 2}, "csv"),
        ({"sheets": [{"name": "A", "columns": COLUMNS, "rows": [["x"]]}]}, "expected 4"),
        ({"sheets": [{"name": "A/B", "columns": COLUMNS, "rows": []}]}, "sheet name"),
        ({"sheets": [{"name": "A", "columns": COLUMNS, "rows": []}] * 2}, "unique"),
    ],
)
def test_invalid_payloads_are_rejected(change, problem):
    with pytest.raises(ValidationError, match=problem):
        _spec(**change)


def test_tool_schema_is_simple_for_every_provider():
    schema = parameters_schema(get_tool("generate_files"))
    text = str(schema)
    assert "anyOf" not in text and "$ref" not in text
    sheet = schema["properties"]["files"]["items"]["properties"]["sheets"]["items"]
    assert sheet["properties"]["rows"]["items"]["items"] == {"type": ["string", "number", "null"]}
    column = sheet["properties"]["columns"]["items"]
    assert column["additionalProperties"] is False
    assert column["properties"]["type"]["enum"] == ["text", "number", "money", "date"]


@pytest.fixture
def ctx(bank, session_a):
    return d.ToolContext(
        bank=bank,
        session=session_a,
        today=NOW.date(),
        files=MemoryFileStore(),
        conversation_id="conv_1",
    )


async def test_tool_stores_files_for_the_session_customer(ctx):
    args = {
        "files": [
            _spec().model_dump(),
            _spec("csv", filename="resumen").model_dump(),
        ]
    }
    result = await get_tool("generate_files").run(ctx, args)
    assert result.ok
    assert [f["filename"] for f in result.data["files"]] == [
        "movimientos junio.xlsx",
        "resumen.csv",
    ]
    assert "file_id" not in str(result.data)  # the model gets no IDs or links
    xlsx, csv_ref = result.files
    assert csv_ref["download_url"] == f"/v1/files/{csv_ref['file_id']}"
    stored = await ctx.files.get(xlsx["file_id"])
    assert stored.customer_id == CUSTOMER_A and stored.conversation_id == "conv_1"
    assert stored.content.startswith(b"PK")  # a zip: xlsx


async def test_tool_reports_bad_arguments_to_the_model(ctx):
    result = await get_tool("generate_files").run(ctx, {"files": []})
    assert not result.ok and result.error_code == "invalid_arguments"


async def test_tool_without_a_store_says_so(bank, session_a):
    ctx = d.ToolContext(bank=bank, session=session_a, today=NOW.date())
    result = await get_tool("generate_files").run(ctx, {"files": [_spec().model_dump()]})
    assert result.error_code == "files_unavailable"


async def test_sqlite_store_round_trip_and_purge(tmp_path):
    store = SqliteFileStore(tmp_path / "db.sqlite")
    [saved] = await service.generate(
        store,
        GenerateFilesArgs(files=[_spec("csv")]),
        customer_id=CUSTOMER_A,
        conversation_id=None,
        now=NOW,
        ttl=timedelta(hours=1),
    )
    loaded = await store.get(saved.file_id)
    assert loaded == saved
    assert await store.purge_expired(NOW + timedelta(hours=2)) == 1
    assert await store.get(saved.file_id) is None
